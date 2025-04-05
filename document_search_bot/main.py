from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
import os
import shutil
import ollama
import pdfplumber
from docx import Document
import pandas as pd
from pptx import Presentation

app = FastAPI()

# Enable CORS to allow frontend requests
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allow all origins
    allow_credentials=True,
    allow_methods=["*"],  # Allow all HTTP methods (GET, POST, DELETE, etc.)
    allow_headers=["*"],  # Allow all headers
)

# Directory to store uploaded documents
UPLOAD_DIR = "uploaded_docs"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@app.post("/upload/")
async def upload_document(file: UploadFile = File(...)):
    """Uploads a document to the server."""
    if file.size > 2 * 1024 * 1024:  # 2MB limit
        raise HTTPException(status_code=400, detail="File size exceeds 2MB limit")
    
    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    
    return {"message": "File uploaded successfully", "filename": file.filename}

@app.get("/query/")
def query_document(question: str):
    """Processes the query and returns an AI-generated response."""
    try:
        document_text = read_uploaded_documents()
        if not document_text:
            raise HTTPException(status_code=400, detail="No documents uploaded or could not be read")
        
        response = ollama.chat(
            model="llama3",
            messages=[
                {"role": "system", "content": "You are an AI assistant that answers questions based on uploaded documents."},
                {"role": "user", "content": f"Document Text: {document_text}\n\nQuestion: {question}"}
            ]
        )

        return {"response": response["message"]["content"]}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

def read_uploaded_documents():
    """Reads uploaded documents and extracts text from supported formats."""
    all_text = ""

    for filename in os.listdir(UPLOAD_DIR):
        file_path = os.path.join(UPLOAD_DIR, filename)
        try:
            if filename.endswith(".txt"):
                with open(file_path, "r", encoding="utf-8") as file:
                    all_text += file.read() + "\n"
            elif filename.endswith(".pdf"):
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            all_text += text + "\n"
            elif filename.endswith(".docx"):
                doc = Document(file_path)
                for para in doc.paragraphs:
                    all_text += para.text + "\n"
            elif filename.endswith((".xls", ".xlsx")):
                df = pd.read_excel(file_path)
                all_text += df.to_string() + "\n"  # Convert table to string
            elif filename.endswith(".pptx"):
                ppt = Presentation(file_path)
                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            all_text += shape.text + "\n"
        except Exception as e:
            print(f"Could not read {filename}: {e}")
            continue

    return all_text.strip()

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000, reload=True)
